"""Assemble the checked-in DeepWeeds run artifacts into a submission folder."""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "checkpoints"
OUT = ROOT / "submission_package"
CLASSES = [
    "Chinee apple", "Lantana", "Parkinsonia", "Parthenium",
    "Prickly acacia", "Rubber vine", "Siam weed", "Snake weed", "Negative",
]
REPO = "https://github.com/HoangDiine/K4-Track4-Day2-Deeplearning-Advance"


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def copy_file(src: Path, dest: Path):
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)


def build_tables():
    results = ARTIFACTS / "results"
    backbones = read_json(results / "backbones.json")
    training = read_json(results / "training.json")
    inference = read_json(results / "inference_summary.json")
    final = read_json(results / "final.json")

    backbone_rows = []
    for row in backbones:
        backbone_rows.append({
            "exp_id": row["exp_id"], "backbone": row["backbone"],
            "pretrained_tag": row["pretrained_tag"], "params_m": row["params_m"],
            "GMAC": row["gmac"], "input_size_px": row["img_size"],
            "best_epoch": row["best_epoch"], "seed": row["seed"],
            "macro_f1_val": row["macro_f1_val"], "top1_val": row["top1_val"],
            "train_seconds_per_epoch": row["train_seconds_per_epoch"],
            "latency_batch1_p95_ms": None, "GPU": row["gpu"], "torch": row["torch"],
            "note": "2 epochs; single screening seed; same T00 training recipe",
        })

    baseline = next(row for row in training if row["exp_id"] == "T00")
    recipe_axis = {
        "T01": ("augmentation", "ColorJitter"),
        "T02": ("loss", "label smoothing 0.1"),
        "T03": ("sampler", "balanced sampler"),
    }
    training_rows = []
    for row in training:
        axis, changed = recipe_axis.get(row["exp_id"], ("baseline", "T00 reference"))
        config = read_json(ARTIFACTS / "runs" / row["exp_id"] / f"seed{row['seed']}" / "config.json")
        training_rows.append({
            "exp_id": row["exp_id"], "backbone": row["backbone"],
            "axis": axis, "change_vs_T00": changed, "seed": row["seed"],
            "macro_f1_val": row["macro_f1_val"], "top1_val": row["top1_val"],
            "delta_macro_f1_vs_T00": row["macro_f1_val"] - baseline["macro_f1_val"],
            "train_seconds_per_epoch": row["train_seconds_per_epoch"],
            "augmentation": config["aug"], "loss": config["loss"],
            "label_smoothing": config["label_smoothing"], "sampler": config["sampler"],
            "note": "3 epochs; one seed; exploratory ablation",
        })

    inference_rows = []
    for row in inference["rows"]:
        inference_rows.append({
            "exp_id": row["exp_id"], "method": row["method"], "model": row["model"],
            "K_views_or_models": row["K (views/models)"],
            "macro_f1_val": row["macro-F1 val"], "top1_val": row["top-1 val"],
            "ece_val": row["ECE val"], "latency_p50_batch1_ms": row["p50 batch-1 (ms)"],
            "latency_p95_batch1_ms": row["p95 batch-1 (ms)"],
            "latency_p99_batch1_ms": row["p99 batch-1 (ms)"],
            "throughput_batch1_img_s": row["throughput batch-1 (img/s)"],
            "latency_p95_batch32_ms": row["p95 batch-32 (ms)"],
            "relative_cost_vs_I00": row["relative cost"], "GPU": row["GPU"],
            "dtype": row["dtype"], "warmup": row["warmup"],
            "measurements": row["measurements"], "temperature": row.get("temperature"),
            "ECE_before": row.get("ECE before"), "ECE_after": row.get("ECE after"),
            "note": row["note"],
        })

    final_by_id = {row["exp_id"]: row for row in final}
    scored = {}
    for exp_id in ("F01", "T00"):
        summary = read_json(results / "eval_out" / f"{exp_id}_summary.json")
        scored[exp_id] = summary
    final_rows = []
    for exp_id in ("F01", "T00"):
        records = [r for r in final if r["exp_id"] == exp_id]
        seed_metrics = pd.read_csv(results / "eval_out" / f"{exp_id}_per_seed.csv").set_index("seed")
        for record in records:
            metrics = seed_metrics.loc[record["seed"]]
            final_rows.append({
                "exp_id": exp_id, "configuration": f"convnext_tiny + T00 + I00; seed {record['seed']}",
                "seed": record["seed"], "macro_f1_val": record["macro_f1_val"],
                "macro_f1_test_saved_eval": metrics["macro_f1"],
                "top1_test_saved_eval": metrics["top1"], "ece_test_saved_eval": metrics["ece"],
                "note": "Scored from saved prediction CSV by eval.py",
            })
        s = scored[exp_id]
        final_rows.append({
            "exp_id": exp_id, "configuration": "convnext_tiny + T00 + I00; mean and sample std over 3 seeds",
            "seed": "mean ± std (n=3)",
            "macro_f1_val": records[0]["macro_f1_val"],
            "macro_f1_test_saved_eval": f"{s['macro_f1']['mean']:.4f} ± {s['macro_f1']['std']:.4f}",
            "top1_test_saved_eval": f"{s['top1']['mean']:.4f} ± {s['top1']['std']:.4f}",
            "ece_test_saved_eval": f"{s['ece']['mean']:.4f} ± {s['ece']['std']:.4f}",
            "note": "Scored from saved predictions by eval.py; F01 and T00 predictions are identical",
        })

    per_seed_rows = []
    for exp_id in ("F01", "T00"):
        frame = pd.read_csv(results / "eval_out" / f"{exp_id}_per_seed.csv")
        for record in frame.to_dict(orient="records"):
            per_seed_rows.append({"exp_id": exp_id, **record})

    per_class_rows = []
    for exp_id in ("F01", "T00"):
        frame = pd.read_csv(results / "eval_out" / f"{exp_id}_per_class.csv")
        for record in frame.to_dict(orient="records"):
            per_class_rows.append({"exp_id": exp_id, **record})

    latency_rows = []
    for row in inference["rows"]:
        latency_rows.append({
            "configuration": row["exp_id"], "GPU": row["GPU"], "dtype": row["dtype"],
            "batch": 1, "BN_fusion": "no", "p50_ms": row["p50 batch-1 (ms)"],
            "p95_ms": row["p95 batch-1 (ms)"], "p99_ms": row["p99 batch-1 (ms)"],
            "images_per_second": row["throughput batch-1 (img/s)"],
            "warmup": row["warmup"], "measurements": row["measurements"],
        })
        latency_rows.append({
            "configuration": row["exp_id"], "GPU": row["GPU"], "dtype": row["dtype"],
            "batch": 32, "BN_fusion": "no", "p50_ms": row["p50 batch-32 (ms)"],
            "p95_ms": row["p95 batch-32 (ms)"], "p99_ms": row["p99 batch-32 (ms)"],
            "images_per_second": row["throughput batch-32 (img/s)"],
            "warmup": row["warmup"], "measurements": row["measurements"],
        })

    candidates = []
    for row in backbone_rows:
        candidates.append({"exp_id": row["exp_id"], "family": "Backbone", "model": row["backbone"],
                          "macro_f1_val": row["macro_f1_val"], "top1_val": row["top1_val"],
                          "p95_ms": None, "params_m": row["params_m"], "GMAC": row["GMAC"],
                          "selection_note": "2-epoch screening; one seed"})
    for row in training_rows:
        candidates.append({"exp_id": row["exp_id"], "family": "Training", "model": row["backbone"],
                           "macro_f1_val": row["macro_f1_val"], "top1_val": row["top1_val"],
                           "p95_ms": None, "params_m": baseline["params_m"], "GMAC": baseline["gmac"],
                           "selection_note": row["change_vs_T00"]})
    for row in inference_rows:
        candidates.append({"exp_id": row["exp_id"], "family": "Inference", "model": row["model"],
                           "macro_f1_val": row["macro_f1_val"], "top1_val": row["top1_val"],
                           "p95_ms": row["latency_p95_batch1_ms"], "params_m": baseline["params_m"],
                           "GMAC": baseline["gmac"], "selection_note": row["method"]})
    top10 = pd.DataFrame(candidates).sort_values("macro_f1_val", ascending=False).head(10).reset_index(drop=True)
    top10.insert(0, "rank", range(1, len(top10) + 1))
    top10["macro_f1_val"] = top10["macro_f1_val"].round(4)
    top10["top1_val"] = top10["top1_val"].round(4)
    return {
        "Summary": top10,
        "Backbones": pd.DataFrame(backbone_rows),
        "Training": pd.DataFrame(training_rows),
        "Inference": pd.DataFrame(inference_rows),
        "Final": pd.DataFrame(final_rows),
        "EvalPerSeed": pd.DataFrame(per_seed_rows),
        "PerClass": pd.DataFrame(per_class_rows),
        "Latency": pd.DataFrame(latency_rows),
    }, backbones, training_rows, inference, final, scored


def write_workbook(tables):
    path = OUT / "results.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet, frame in tables.items():
            frame.to_excel(writer, sheet_name=sheet, index=False)
        workbook = writer.book
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17365D")
            for column in sheet.columns:
                letter = get_column_letter(column[0].column)
                width = min(52, max(12, max(len(str(cell.value or "")) for cell in column[:60]) + 2))
                sheet.column_dimensions[letter].width = width


def report_text(backbones, training, inference, final, scored, grade):
    b = max(backbones, key=lambda r: r["macro_f1_val"])
    baseline = next(row for row in training if row["exp_id"] == "T00")
    recipes = {row["exp_id"]: row for row in training}
    infer_rows = {row["exp_id"]: row for row in inference["rows"]}
    best_inf = max(inference["rows"], key=lambda r: r["macro-F1 val"])
    i00, i01, i02, i07 = (infer_rows[key] for key in ("I00", "I01", "I02", "I07"))
    f, t = scored["F01"], scored["T00"]
    f_class = pd.read_csv(ARTIFACTS / "results/eval_out/F01_per_class.csv")
    f_conf = pd.read_csv(ARTIFACTS / "results/eval_out/F01_confusion_sum.csv", index_col=0)
    top_errors = []
    for true_name in f_conf.index:
        for pred_name, count in f_conf.loc[true_name].items():
            truth = true_name.removeprefix("true_")
            prediction = pred_name.removeprefix("pred_")
            if truth != prediction:
                top_errors.append((int(count), truth, prediction))
    top_errors.sort(reverse=True)
    top_errors = top_errors[:6]
    per_class = "\n".join(
        f"| {r['class']} | {int(r['support'])} | {r['precision_mean']:.3f} | {r['recall_mean']:.3f} | {r['f1_mean']:.3f} |"
        for r in f_class.to_dict(orient="records")
    )
    conf_table = "\n".join(f"| {n} | {b['macro_f1_val']:.4f} | {b['top1_val']:.4f} | {b['params_m']:.2f} | {b['gmac']:.2f} |" for n, b in [(r["backbone"], r) for r in backbones])
    train_table = "\n".join(
        f"| {r['exp_id']} | {r['augmentation']} / loss={r['loss']} / sampler={r['sampler']} | {r['macro_f1_val']:.4f} | {r['macro_f1_val']-baseline['macro_f1_val']:+.4f} | {r['train_seconds_per_epoch']:.1f} |"
        for r in training
    )
    inference_table = "\n".join(
        f"| {r['exp_id']} | {r['method']} | {r['macro-F1 val']:.4f} | {r['ECE val']:.4f} | {r['p95 batch-1 (ms)']:.2f} | {r['relative cost']:.2f}× |"
        for r in inference["rows"] if r["exp_id"] in {"I00", "I01", "I02", "I04", "I07"}
    )
    final_rows = []
    for exp_id, summary in (("F01", f), ("T00", t)):
        final_rows.append(
            f"| {exp_id} | {summary['top1']['mean']:.4f} ± {summary['top1']['std']:.4f} | "
            f"{summary['macro_f1']['mean']:.4f} ± {summary['macro_f1']['std']:.4f} | "
            f"{summary['ece']['mean']:.4f} ± {summary['ece']['std']:.4f} |"
        )
    error_lines = "\n".join(f"- {truth} → {pred}: {n} occurrences across the three F01 test seeds." for n, truth, pred in top_errors)
    confusion_header = "| Thật \\ Dự đoán | " + " | ".join(c.removeprefix("pred_") for c in f_conf.columns) + " |"
    confusion_sep = "|---|" + "---:|" * len(f_conf.columns)
    confusion_rows = "\n".join(
        "| " + true_name.removeprefix("true_") + " | " + " | ".join(str(int(value)) for value in f_conf.loc[true_name]) + " |"
        for true_name in f_conf.index
    )
    graded_points = sum(item["points"] or 0 for item in grade["items"])
    graded_total = sum(item["max"] for item in grade["items"] if item["points"] is not None)
    return f'''# DeepWeeds — báo cáo thí nghiệm

## 1. Tóm tắt

- Phân loại 9 lớp DeepWeeds, dùng fold 0 với 17.509 ảnh; các tập train/val/test lần lượt có 10.501/3.501/3.507 ảnh.
- Đã ghi nhận 5 backbone, 3 biến thể công thức ngoài baseline, 7 phương pháp suy luận và vòng chung kết 3 seed cho F01/T00.
- Kết quả `eval.py` trên dự đoán đã lưu: F01 đạt macro-F1 **{f['macro_f1']['mean']:.4f} ± {f['macro_f1']['std']:.4f}**, top-1 **{f['top1']['mean']:.4f} ± {f['top1']['std']:.4f}**; T00 cho kết quả giống hệt.
- Không có cải thiện đo được của F01 so với T00. Hai nhóm dự đoán trùng nhau theo seed; F01 đã chạy với cùng backbone, công thức T00 và suy luận I00.
- I01 có validation macro-F1 cao nhất ({best_inf['macro-F1 val']:.4f}), nhưng phép suy luận đó **không được dùng ở vòng test cuối**; không suy rộng điểm validation sang test.
- Theo `eval.py grade` trên các dự đoán đã lưu: {graded_points}/{graded_total} điểm trong tiêu chí đã chấm; I4a chưa chấm do không có dự đoán test sau temperature scaling.

## 2. Dữ liệu và thiết lập

Nguồn là DeepWeeds, fold 0 nguyên bản. Audit ghi nhận 0 ảnh giao giữa các tập, hợp đủ 17.509 ảnh và không thiếu tệp. Negative chiếm 5.463/10.501 ảnh train (52,0%), vì vậy top-1 không đủ để đánh giá riêng chất lượng các lớp hiếm.

![Phân bố lớp train](report_assets/class_distribution.png)

Các phép sàng lọc backbone dùng 2 epoch và một seed; ablation dùng 3 epoch và một seed; chung kết dùng 3 seed. Mọi chọn lựa backbone/công thức/phương pháp suy luận dựa trên validation. GPU ghi trong log là Tesla T4, PyTorch 2.11.0+cu130; tag trọng số từng backbone nằm trong workbook. Batch size 32, ảnh 224 px, AMP bật. Giới hạn thời gian khiến thí nghiệm ngắn hơn cấu hình khuyến nghị trong rubric.

## 3. So sánh backbone

| Backbone | Macro-F1 val | Top-1 val | Params (M) | GMAC |
|---|---:|---:|---:|---:|
{conf_table}

`{b['backbone']}` là kết quả validation tốt nhất trong nhóm sàng lọc ở {b['macro_f1_val']:.4f}; được chọn làm backbone tiếp tục. Đây là sàng lọc ngắn, mỗi kiến trúc chỉ một seed và hai epoch, nên không đủ bằng chứng để kết luận ưu thế tổng quát giữa các kiến trúc.

## 4. Công thức huấn luyện

| ID | Thay đổi so với T00 | Macro-F1 val | Δ vs T00 | Giây/epoch |
|---|---|---:|---:|---:|
{train_table}

 T01 (ColorJitter) giảm macro-F1 validation; T02 label smoothing gần baseline về macro-F1 nhưng ECE validation tăng mạnh; T03 balanced sampler không cho thấy mức tăng nhất quán trong một seed. Chênh lệch nhỏ không thể phân biệt với nhiễu do chỉ có một lần chạy ở mỗi ablation. Vì vậy công thức cuối không được chứng minh là tốt hơn T00.

## 5. Phương pháp suy luận

| ID | Phương pháp | Macro-F1 val | ECE val | p95 batch 1 (ms) | Chi phí tương đối |
|---|---|---:|---:|---:|---:|
{inference_table}

I01 (lật ngang, hai view) tăng macro-F1 validation từ {i00['macro-F1 val']:.4f} lên {i01['macro-F1 val']:.4f}, còn p95 tăng từ {i00['p95 batch-1 (ms)']:.2f} lên {i01['p95 batch-1 (ms)']:.2f} ms. I02 dùng năm crop, tốn p95 {i02['p95 batch-1 (ms)']:.2f} ms mà không vượt I01. I07 khớp T={i07['temperature']:.4f} chỉ bằng validation; ECE validation giảm từ {i07['ECE before']:.4f} xuống {i07['ECE after']:.4f}, Top-1 không đổi.

**Giới hạn quan trọng:** test cuối là suy luận I00 (một view, không temperature scaling). Notebook đã chọn I01 trên validation nhưng không chuyển lựa chọn này vào F01 khi tạo dự đoán test. Vì quy tắc yêu cầu test một lần mỗi seed và dự đoán đã được tạo, báo cáo giữ số hiện có, không tạo lần đánh giá test mới. Do đó không có con số test cho I01 hoặc temperature scaling.

## 6. Chung kết và lỗi

| Cấu hình | Top-1 test (mean ± std) | Macro-F1 test (mean ± std) | ECE test (mean ± std) |
|---|---:|---:|---:|
{chr(10).join(final_rows)}

Các số lấy từ `eval.py` chạy trên CSV dự đoán đã lưu, không lấy từ trường `test_metrics` trong log train vì các giá trị đó không khớp với bộ chấm chính thức. T00 và F01 có cùng dự đoán ở cả ba seed; chênh lệch macro-F1 là 0.0000, nhỏ hơn độ lệch chuẩn mẫu 0.0025.

### Precision / recall / F1 theo lớp (F01)

| Lớp | N test | Precision | Recall | F1 |
|---|---:|---:|---:|---:|
{per_class}

Recall thấp nhất là Snake weed ({f_class.loc[f_class['class']=='Snake weed','recall_mean'].iloc[0]:.3f}), kế đến Chinee apple ({f_class.loc[f_class['class']=='Chinee apple','recall_mean'].iloc[0]:.3f}). Các nhầm lẫn thường gặp nhất trong ma trận cộng gộp ba seed:

{error_lines}

Ma trận dưới đây là số lượng đúng / sai cộng qua 3 seed (hàng là nhãn thật, cột là nhãn dự đoán; tổng N=10.521):

{confusion_header}
{confusion_sep}
{confusion_rows}

Negative là lớp lớn và nhiều loài cỏ bị đoán thành Negative; ma trận cũng ghi nhận nhầm lẫn hai chiều Chinee apple ↔ Snake weed. Ảnh ví dụ trong `report_assets/error_examples.png` được lấy từ các lỗi seed 0 để xem thủ công. Đây là gợi ý trực quan, không thay thế đánh giá định lượng.

## 7. Khuyến nghị và hạn chế

- Với kết quả test đã có, dùng T00 + ConvNeXt-Tiny + I00 làm cấu hình chuẩn: macro-F1 0.9618 ± 0.0025, p95 batch-1 7.10 ms trên Tesla T4. Số latency là trên T4, không đại diện cho phần cứng robot.
- I01 là ứng viên nếu cần tối đa hóa điểm validation và ngân sách cho phép p95 khoảng 33 ms; chưa có điểm test cho phương pháp này.
- Bằng chứng hiện tại không cho thấy F01 cải thiện chất lượng so với baseline. Không kết luận các biến thể nhỏ hơn nhiễu là tốt hơn.
- Hạn chế: một fold, screening/ablation một seed, chỉ 2–3 epoch; một GPU; dự đoán F01/T00 trùng nhau; F01 chưa dùng I01; chưa có hiệu chuẩn test sau temperature scaling. Cần ghi nhận đây là giới hạn run đã hoàn tất, không tự điền kết quả còn thiếu.

## 8. Tái lập và tệp đính kèm

Workbook `results.xlsx` chứa Summary, Backbones, Training, Inference, Final, EvalPerSeed, PerClass và Latency. Code tái lập nằm trong `code/`, đồ thị trong `curves/`, CSV dự đoán chung kết/validation trong `predictions/`, còn log JSON/CSV gốc đã dùng để tạo bảng nằm trong `raw_results/`. Checkpoint không được đóng gói vì mỗi checkpoint vượt 16 MB; các đường dẫn và file `best.pth` gốc nằm dưới thư mục ignored `checkpoints/runs/` của workspace.
'''


def main():
    if not ARTIFACTS.is_dir():
        raise FileNotFoundError(f"Missing checkpoint artifacts: {ARTIFACTS}")
    tables, backbones, training, inference, final, scored = build_tables()
    (OUT / "code").mkdir(parents=True, exist_ok=True)
    for name in ("benchmark.py", "dataset.py", "inference.py", "losses.py", "model.py", "train.py"):
        copy_file(ROOT / "starter" / name, OUT / "code" / name)
    copy_file(ROOT / "eval.py", OUT / "code" / "eval.py")
    copy_file(ROOT / "scripts" / "audit_deepweeds.py", OUT / "code" / "audit_deepweeds.py")
    copy_file(ROOT / "starter" / "lab_day2.ipynb", OUT / "lab_day2.ipynb")
    for source in (ARTIFACTS / "curves").glob("*.png"):
        copy_file(source, OUT / "curves" / source.name)
    for exp_id in ("F01", "T00"):
        for seed in range(3):
            for split in ("test", "val"):
                source = ARTIFACTS / "predictions" / f"{exp_id}_seed{seed}_{split}.csv"
                if source.is_file():
                    copy_file(source, OUT / "predictions" / source.name)
    for source in (
        ARTIFACTS / "results/backbones.json", ARTIFACTS / "results/training.json",
        ARTIFACTS / "results/final.json", ARTIFACTS / "results/inference_summary.json",
        ARTIFACTS / "results/eval_out/F01_summary.json", ARTIFACTS / "results/eval_out/T00_summary.json",
        ARTIFACTS / "results/eval_out/F01_per_seed.csv", ARTIFACTS / "results/eval_out/T00_per_seed.csv",
        ARTIFACTS / "results/eval_out/F01_per_class.csv", ARTIFACTS / "results/eval_out/T00_per_class.csv",
        ARTIFACTS / "results/eval_out/F01_confusion_sum.csv", ARTIFACTS / "results/eval_out/T00_confusion_sum.csv",
        ROOT / "eda/split_audit.json",
    ):
        copy_file(source, OUT / "raw_results" / source.name)
    for source in (ARTIFACTS / "runs").glob("*/seed*"):
        for pattern in ("config.json", "result.json", "history.csv"):
            f = source / pattern
            if f.exists():
                copy_file(f, OUT / "raw_results/runs" / source.parent.name / source.name / f.name)

    realtime_p95 = next(row["p95 batch-1 (ms)"] for row in inference["rows"] if row["exp_id"] == "I00")
    grade_command = [
        sys.executable, "-X", "utf8", str(ROOT / "eval.py"), "grade",
        "--final", str(OUT / "predictions/F01_seed*_test.csv"),
        "--baseline", str(OUT / "predictions/T00_seed*_test.csv"),
        "--final-val", str(OUT / "predictions/F01_seed*_val.csv"),
        "--val-csv", str(ROOT / "data/labels/val_subset0.csv"),
        "--test-csv", str(ROOT / "data/labels/test_subset0.csv"),
        "--labels", str(ROOT / "data/labels/labels.csv"),
        "--latency-p95-ms", str(realtime_p95),
        "--out", str(OUT / "raw_results/grade_I_regraded"),
    ]
    subprocess.run(grade_command, cwd=ROOT, check=True)
    grade = read_json(OUT / "raw_results/grade_I_regraded/grade_I.json")
    copy_file(ROOT / "eda/class_distribution.png", OUT / "report_assets/class_distribution.png")

    # Choose a few real seed-0 errors from the difficult Chinee apple / Snake weed pair.
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from PIL import Image
    pred = pd.read_csv(ARTIFACTS / "predictions/F01_seed0_test.csv")
    errors = pred[(pred.y_true != pred.y_pred) & pred.y_true.isin([0, 7]) & pred.y_pred.isin([0, 7])].head(4)
    fig, axes = plt.subplots(1, max(1, len(errors)), figsize=(12, 3.6))
    if len(errors) == 0:
        axes = [axes]
    for ax, (_, row) in zip(axes, errors.iterrows()):
        image_path = ROOT / "images" / row.Filename
        ax.imshow(Image.open(image_path).convert("RGB"))
        ax.set_title(f"true={CLASSES[int(row.y_true)]}\npred={CLASSES[int(row.y_pred)]}")
        ax.axis("off")
    fig.tight_layout()
    (OUT / "report_assets").mkdir(parents=True, exist_ok=True)
    fig.savefig(OUT / "report_assets/error_examples.png", dpi=160, bbox_inches="tight")
    plt.close(fig)

    write_workbook(tables)
    (OUT / "report.md").write_text(report_text(backbones, training, inference, final, scored, grade), encoding="utf-8")
    (OUT / "README.md").write_text(f'''# DeepWeeds Lab — submission package

## Contents

- `results.xlsx`: comparison tables and saved `eval.py` scores.
- `report.md`: results, analysis, limits, and recommendation.
- `curves/`: training curves for every recorded B/T/F run.
- `code/`: source modules and evaluation code.
- `predictions/`: final and validation CSVs for seeds 0–2.
- `raw_results/`: compact JSON/CSV source records; no model weights.
- `lab_day2.ipynb`: executed Colab notebook used for this run.

## Re-run

Open the notebook in [Google Colab](https://colab.research.google.com/github/HoangDiine/K4-Track4-Day2-Deeplearning-Advance/blob/main/starter/lab_day2.ipynb), select a GPU runtime, and run cells top to bottom. The notebook mounts Drive and stores generated runs there. To reproduce this exact local revision, upload the included notebook and `code/` folder to Colab (the public branch may not yet contain local fixes). Python package versions and the recorded hardware are in the notebook and workbook. The saved run used seed 0 for screening and seeds 0, 1, 2 for the final/baseline.

For local scoring of saved predictions, from the repository root run `python eval.py score` and `python eval.py grade`; see the notebook's scoring cells for complete arguments.

The final test outputs used I00. I01 was evaluated on validation only and must not be presented as a test result.
''', encoding="utf-8")
    print(f"Created {OUT}")
    print(f"Workbook sheets: {list(tables)}")
    print(f"Curves: {len(list((OUT/'curves').glob('*.png')))}; predictions: {len(list((OUT/'predictions').glob('*.csv')))}")


if __name__ == "__main__":
    main()
